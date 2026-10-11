#!/usr/bin/env python3
"""
Импорт сметы из Excel. stdin: файл .xlsx (байты). stdout: JSON {title, lines, warnings}.
Ищет строку заголовков (наименование / ед. / кол-во / цена / сумма), разделы — строки без чисел,
состав позиции — строки ячейки после первой («– сетка 1,73 м…»).
"""
import io
import json
import re
import sys

from openpyxl import load_workbook

SKIP = re.compile(r"^\s*(итого|всего|ндс|в т\.?\s*ч|сметная стоимость|получено|оплачено|к оплате|сумма прописью|подрядчик|заказчик|исполнитель|примечание)", re.I)


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace("₽", "").replace("руб.", "").replace("р.", "").replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None


def find_header(ws):
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=False):
        texts = {c.column: str(c.value).strip().lower() for c in r if c.value is not None and isinstance(c.value, str)}
        joined = " ".join(texts.values())
        if re.search(r"наименован|работ[аы]?\b|позици|товар|номенклатур", joined) and re.search(r"кол|цена|сумм|стоим", joined):
            return r[0].row, texts
    return None, {}


def map_cols(texts, pricelist=False):
    m = {}
    if pricelist:
        prices = [c for c, t in sorted(texts.items()) if "цена" in t or "стоим" in t or "руб" in t]
        want = "опт" if "--opt" in sys.argv else "розн"
        pref = [c for c in prices if want in texts[c]]
        if prices:
            m["price"] = (pref or prices)[0]
    for col, t in sorted(texts.items()):
        if col == m.get("price"):
            continue
        if "наименован" in t or "товар" in t or "номенклатур" in t or ("работ" in t and "name" not in m and "цена" not in t and "стоим" not in t):
            m.setdefault("name", col)
        elif re.search(r"\bед", t):
            m.setdefault("unit", col)
        elif "кол" in t:
            m.setdefault("qty", col)
        elif "матер" in t and not pricelist:
            m.setdefault("mat", col)
        elif "работ" in t and ("цена" in t or "стоим" in t or "р." in t or "руб" in t) and not pricelist:
            m.setdefault("work", col)
        elif ("цена" in t or "за ед" in t) and not pricelist:
            m.setdefault("price", col)
        elif ("сумм" in t or "стоим" in t or "всего" in t) and not pricelist:
            m.setdefault("sum", col)
    return m


def parse(data, pricelist=False):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    best = None
    for ws in wb.worksheets:
        row, texts = find_header(ws)
        if row:
            best = (ws, row, texts)
            break
    if not best:
        raise SystemExit("Не нашли таблицу: нужна строка заголовков с колонками «Наименование», «Кол-во», «Цена» или «Сумма»")
    ws, hrow, texts = best
    cols = map_cols(texts, pricelist)
    if "name" not in cols:
        raise SystemExit("Не нашли колонку «Наименование»")
    title = ""
    for r in ws.iter_rows(min_row=1, max_row=hrow - 1, values_only=True):
        for v in r:
            if isinstance(v, str) and len(v.strip()) > 3 and not SKIP.match(v) and not re.match(r"^(от|№)\s", v.strip(), re.I):
                title = title or v.strip()
    lines, warnings, section = [], [], ""
    get = lambda r, k: r[cols[k] - 1] if k in cols and cols[k] - 1 < len(r) else None
    for r in ws.iter_rows(min_row=hrow + 1, max_row=min(ws.max_row, hrow + 3000), values_only=True):
        raw = get(r, "name")
        if raw is None or str(raw).strip() == "":
            # раздел в объединённой ячейке левее колонки наименования
            texts = [str(v).strip() for v in r if isinstance(v, str) and v.strip()]
            nums = [v for v in r if isinstance(v, (int, float))]
            if len(texts) == 1 and not nums and not SKIP.match(texts[0]):
                t0 = texts[0][:100].strip(" :")
                section = t0.capitalize() if t0.isupper() else t0
            continue
        text = str(raw).strip()
        if re.fullmatch(r"[\d\s.]+", text):  # строка с номерами колонок «1 2 3…»
            continue
        if SKIP.match(text):
            continue
        q, p, w, m, s = (to_num(get(r, k)) for k in ("qty", "price", "work", "mat", "sum"))
        if q is None and p is None and w is None and m is None and s is None:
            if text.lower() not in ("работа", "работы", "материалы"):
                section = text[:100].strip(" :").capitalize() if text.isupper() else text[:100].strip(" :")
            continue
        parts = [x.strip() for x in re.split(r"\n|(?<=[,;:])\s*(?=[-–•]\s)", text) if x.strip()]
        name = re.sub(r",?\s*в\s+составе\s*:?\s*$", "", parts[0], flags=re.I).strip(" ,:")
        comp = [re.sub(r"^[-–•]\s*", "", x).strip(" ;,") for x in parts[1:]]
        qty = q if q is not None else 1
        if w is not None or m is not None:
            work, mat = w or 0, m or 0
        else:
            price = p if p is not None else (s / qty if s is not None and qty else 0)
            work, mat = price, 0
        if not (work or mat):
            warnings.append(f"Без цены: {name[:60]}")
        unit = str(get(r, "unit") or "шт").strip()[:30] or "шт"
        lines.append({"name": name[:300] or "Позиция", "unit": unit, "qty": round(qty, 4), "price": round(work, 2), "matPrice": round(mat, 2) if mat else None,
                      "category": section or None, "composition": "\n".join(comp)[:3000] or None, "type": "work"})
    if not lines:
        raise SystemExit("В таблице не нашлось позиций с количеством или ценой")
    return {"title": (title or ws.title or "Смета из Excel")[:200], "lines": lines[:2000], "warnings": warnings[:20]}


if __name__ == "__main__":
    try:
        print(json.dumps(parse(sys.stdin.buffer.read(), pricelist='--pricelist' in sys.argv), ensure_ascii=False))
    except SystemExit as e:
        sys.stderr.write(str(e))
        sys.exit(2)
    except Exception as e:  # битый файл и т.п.
        sys.stderr.write(f"Не удалось прочитать файл: {e}")
        sys.exit(2)
