#!/usr/bin/env python3
"""Экспорт в Excel (openpyxl). Читает JSON из stdin, пишет XLSX в stdout.

kind = "estimate": листы «Смета» (для заказчика, с формулами) и «Для себя» (себестоимость, прибыль).
kind = "project":  листы «Сводка» и «Операции» по объекту.
"""
import io
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONEY = '#,##0.00 "₽"'
QTY = '#,##0.####'
HEAD_FILL = PatternFill("solid", fgColor="2B3440")
HEAD_FONT = Font(bold=True, color="FFFFFF")
CAT_FILL = PatternFill("solid", fgColor="E8ECF1")
TOTAL_FILL = PatternFill("solid", fgColor="FFE7A3")
THIN = Side(style="thin", color="C9CFD8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def head(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def kv(ws, row, label, value, fmt=MONEY, bold=False, fill=None, col=1):
    a = ws.cell(row=row, column=col, value=label)
    b = ws.cell(row=row, column=col + 1, value=value)
    if fmt and not isinstance(value, str):
        b.number_format = fmt
    if bold:
        a.font = b.font = Font(bold=True)
    if fill:
        a.fill = b.fill = fill
    return b


def title_block(ws, d, title):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    info = [("Исполнитель", d.get("company")), ("Контакты", d.get("contacts")), ("Заказчик", d.get("customer")),
            ("Адрес", d.get("address")), ("Дата", d.get("date"))]
    r = 2
    for k, v in info:
        if v:
            ws.cell(row=r, column=1, value=k).font = Font(color="666666")
            ws.cell(row=r, column=2, value=v)
            r += 1
    return r + 1


def estimate_sheet(wb, d):
    ws = wb.active
    ws.title = "Смета"
    r = title_block(ws, d, f"Смета № {d.get('number', '')} — {d.get('title', '')}")
    head(ws, r, ["№", "Наименование", "Ед.", "Кол-во", "Цена", "Сумма"], [6, 52, 9, 12, 15, 17])
    ws.freeze_panes = ws.cell(row=r + 1, column=1)
    r += 1
    first = r
    cat = object()
    n = 0
    for l in d["lines"]:
        if d.get("grouped") and l.get("category") != cat:
            cat = l.get("category")
            c = ws.cell(row=r, column=1, value=cat or "Прочее")
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            c.font, c.fill = Font(bold=True), CAT_FILL
            r += 1
        n += 1
        vals = [n, l["name"] + (f" ({l['note']})" if l.get("note") else ""), l["unit"], l["qty"], l["price"], f"=D{r}*E{r}"]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i == 2:
                c.alignment = WRAP
        ws.cell(row=r, column=4).number_format = QTY
        ws.cell(row=r, column=5).number_format = MONEY
        ws.cell(row=r, column=6).number_format = MONEY
        r += 1
    last = r - 1
    r += 1
    t = d["totals"]
    kv(ws, r, "Итого по позициям", f"=SUM(F{first}:F{last})" if last >= first else 0, col=5, bold=True); r += 1
    for label, key in [("Накладные расходы", "overhead"), ("Доставка", "delivery"), ("Прочие расходы", "extraCosts")]:
        if t.get(key):
            kv(ws, r, label, t[key], col=5); r += 1
    if t.get("discount"):
        kv(ws, r, f"Скидка {t.get('discountPct', 0):g}%", -t["discount"], col=5); r += 1
    if t.get("tax"):
        kv(ws, r, f"НДС/налог {t.get('taxPct', 0):g}%", t["tax"], col=5); r += 1
    kv(ws, r, "ИТОГО", t["total"], col=5, bold=True, fill=TOTAL_FILL); r += 2
    if d.get("paymentTerms"):
        ws.cell(row=r, column=1, value="Условия оплаты:").font = Font(bold=True)
        ws.cell(row=r, column=2, value=d["paymentTerms"]).alignment = WRAP; r += 1
    if d.get("notes"):
        ws.cell(row=r, column=1, value="Примечания:").font = Font(bold=True)
        ws.cell(row=r, column=2, value=d["notes"]).alignment = WRAP
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0

    # Лист «Для себя» — себестоимость и маржа по строкам (только в версии для себя)
    if not d.get("own"):
        return
    own = wb.create_sheet("Для себя")
    own["A1"] = "Себестоимость и прибыль (не для заказчика)"
    own["A1"].font = Font(bold=True, size=13)
    head(own, 3, ["Наименование", "Ед.", "Кол-во", "Цена", "Себест. ед.", "Выручка", "Себестоимость", "Прибыль", "Маржа"],
         [46, 8, 11, 14, 14, 16, 16, 16, 10])
    own.freeze_panes = "A4"
    r = 4
    for l in d["lines"]:
        cp = l.get("costPrice")
        vals = [l["name"], l["unit"], l["qty"], l["price"], cp if cp is not None else None,
                f"=C{r}*D{r}", f"=IF(E{r}=\"\",\"\",C{r}*E{r})", f"=IF(E{r}=\"\",\"\",F{r}-G{r})",
                f"=IF(OR(E{r}=\"\",F{r}=0),\"\",H{r}/F{r})"]
        for i, v in enumerate(vals, 1):
            c = own.cell(row=r, column=i, value=v)
            c.border = BORDER
            c.number_format = {3: QTY, 9: "0.0%"}.get(i, MONEY if i >= 4 else "General")
        r += 1
    end = r - 1
    r += 1
    if end >= 4:
        kv(own, r, "Выручка по позициям", f"=SUM(F4:F{end})", col=5, bold=True); r += 1
        kv(own, r, "Себестоимость (где указана)", f"=SUM(G4:G{end})", col=5); r += 1
        kv(own, r, "Прибыль по позициям", f"=SUM(H4:H{end})", col=5, bold=True, fill=TOTAL_FILL); r += 2
    s = d.get("own") if isinstance(d.get("own"), dict) else {}
    for label, key in [("Итого заказчику", "total"), ("Прямые затраты (оценка)", "directCost"), ("Плановая прибыль", "profit")]:
        if key in s:
            kv(own, r, label, s[key], col=5, bold=key == "profit"); r += 1
    if "marginPct" in s:
        kv(own, r, "Маржа", (s["marginPct"] or 0) / 100, fmt="0.0%", col=5)


def project_book(wb, d):
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = f"Объект: {d.get('name', '')}"
    ws["A1"].font = Font(bold=True, size=14)
    r = 2
    for k in ("customer", "address", "status"):
        if d.get(k):
            ws.cell(row=r, column=1, value={"customer": "Заказчик", "address": "Адрес", "status": "Статус"}[k]).font = Font(color="666666")
            ws.cell(row=r, column=2, value=d[k]); r += 1
    r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20
    f = d["finance"]
    rows = [("Сумма договора", "contractTotal"), ("Получено от заказчика", "received"), ("Долг заказчика", "customerDebt"),
            ("Расходы начислено", "expensesAccrued"), ("Расходы оплачено", "expensesPaid"), ("Не оплачено поставщикам/рабочим", "unpaidObligations"),
            ("Деньги на руках", "cashBalance"), ("Плановая себестоимость", "plannedCost"), ("Плановая прибыль", "plannedProfit"),
            ("Текущая прибыль", "currentProfit"), ("Прогноз прибыли", "forecastProfit")]
    for label, key in rows:
        if key in f:
            kv(ws, r, label, f[key], bold=key in ("forecastProfit", "currentProfit"),
               fill=TOTAL_FILL if key == "forecastProfit" else None); r += 1
    if f.get("forecastMarginPct") is not None:
        kv(ws, r, "Прогноз маржи", (f["forecastMarginPct"] or 0) / 100, fmt="0.0%")

    ops = wb.create_sheet("Операции")
    head(ops, 1, ["Дата", "Тип", "Категория", "Название", "Контрагент", "Кол-во", "Ед.", "Сумма", "Оплачено", "Примечание"],
         [12, 11, 22, 36, 22, 10, 7, 15, 15, 30])
    ops.freeze_panes = "A2"
    ops.auto_filter.ref = f"A1:J{max(len(d['ops']) + 1, 2)}"
    r = 2
    for o in d["ops"]:
        vals = [o.get("date"), "Доход" if o["kind"] == "income" else "Расход", o.get("category"), o.get("title"), o.get("counterparty"),
                o.get("qty"), o.get("unit"), o["amount"], o.get("paidAmount") if o["kind"] == "expense" else o["amount"], o.get("note")]
        for i, v in enumerate(vals, 1):
            c = ops.cell(row=r, column=i, value=v)
            c.border = BORDER
            if i in (8, 9):
                c.number_format = MONEY
            if i == 6:
                c.number_format = QTY
        r += 1
    last = r - 1
    r += 1
    if last >= 2:
        kv(ops, r, "Доходы", f'=SUMIF(B2:B{last},"Доход",H2:H{last})', col=7, bold=True); r += 1
        kv(ops, r, "Расходы", f'=SUMIF(B2:B{last},"Расход",H2:H{last})', col=7, bold=True); r += 1
        kv(ops, r, "Разница", f"=H{r-2}-H{r-1}", col=7, bold=True, fill=TOTAL_FILL)


def main():
    d = json.load(sys.stdin)
    wb = Workbook()
    if d.get("kind") == "project":
        project_book(wb, d)
    else:
        estimate_sheet(wb, d)
    buf = io.BytesIO()
    wb.save(buf)
    sys.stdout.buffer.write(buf.getvalue())


if __name__ == "__main__":
    main()
